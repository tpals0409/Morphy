# /// script
# dependencies = ["python-docx", "openpyxl", "python-pptx"]
# ///
"""Office → PDF test corpus (PRD 17 문서 품질). Usage: uv run scripts/make_corpus.py <out_dir>"""
import os, shutil, struct, subprocess, sys, zlib
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PI

out = Path(sys.argv[1]); out.mkdir(parents=True, exist_ok=True)

def png(path, w=120, h=80):  # blue/orange gradient, stdlib only
    rows = b"".join(b"\0" + bytes(c for x in range(w) for c in (x * 2, 100, 255 - x * 2)) for _ in range(h))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))
img = out / "_img.png"; png(img)

d = Document()
d.sections[0].header.paragraphs[0].text = "머리말: 모르피 테스트 보고서"
d.sections[0].footer.paragraphs[0].text = "꼬리말: 대외비"
d.add_heading("분기 보고서 — 한국어 글꼴 검증", 0)
d.add_paragraph("가나다라마바사 아자차카타파하. 다람쥐 헌 쳇바퀴에 타고파. " * 6)
d.add_picture(str(img), width=Inches(3))
t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "번호", "항목", "금액"
for i in range(1, 61):
    r = t.add_row().cells; r[0].text, r[1].text, r[2].text = str(i), f"긴 표 항목 {i}", f"{i * 1000:,}원"
d.add_page_break(); d.add_paragraph("마지막 페이지 문단 END-OF-DOC")
d.save(out / "보고서 한글.docx")

wb = Workbook(); ws = wb.active; ws.title = "매출"
ws.append(["월", "매출", "비용", "이익"])
for m in range(1, 13): ws.append([f"{m}월", m * 100, m * 60, f"=B{m+1}-C{m+1}"])
ws["F1"] = "인쇄 영역 밖 — 나오면 안 됨"
ws.print_area = "A1:D13"; ws.page_setup.orientation = "landscape"
ws2 = wb.create_sheet("두번째 시트"); ws2["A1"] = "두 번째 시트 내용 SHEET-TWO"
wb.save(out / "매출 표.xlsx")

def deck(name, w, h):
    p = Presentation(); p.slide_width, p.slide_height = PI(w), PI(h)
    s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = "발표 제목"; s.placeholders[1].text = f"{name} 부제목"
    s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "도형 슬라이드"
    s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(1), PI(2), PI(3), PI(1.5)).text = "둥근 사각형"
    s.shapes.add_picture(str(img), PI(5), PI(2), PI(2))
    s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "차트 슬라이드"
    cd = CategoryChartData(); cd.categories = ["1분기", "2분기", "3분기"]; cd.add_series("매출", (3, 5, 4))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PI(1), PI(1.8), PI(6), PI(3.5), cd)
    p.save(out / name)
deck("발표 16x9.pptx", 13.333, 7.5)
deck("발표 4x3.pptx", 10, 7.5)
wb = Workbook(); wb.active.append(["이름", "금액"]); wb.active.append(["홍길동", 1000]); wb.save(out / "단일 시트.xlsx")

csv = "이름,금액,메모\n홍길동,1000,\"쉼표, 포함\"\n김철수,2500,일반\n"
(out / "명단 utf8.csv").write_text(csv, encoding="utf-8")
(out / "명단 cp949.csv").write_bytes(csv.encode("cp949"))
(out / "메모 cp949.txt").write_bytes("한글 텍스트 파일입니다.\n두 번째 줄.\n".encode("cp949"))
(out / "README 한글.md").write_text("# 마크다운 제목\n\n- 목록 **굵게**\n- 두 번째\n\n| 이름 | 값 |\n|---|---|\n| 가 | 1 |\n\n```\ncode 블록\n```\n", encoding="utf-8")

# Legacy/ODF inputs, made by LibreOffice itself from the files above.
so = os.environ.get("MORPHY_ENGINE_PATH") or shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
profile = (out.resolve() / ".lo-corpus-profile").as_uri()
for src, targets in [("보고서 한글.docx", ["doc", "odt", "rtf"]), ("매출 표.xlsx", ["xls", "ods"]), ("발표 16x9.pptx", ["ppt", "odp"])]:
    for t in targets:
        subprocess.run([so, "-env:UserInstallation=" + profile, "--headless", "--norestore", "--convert-to", t, "--outdir", str(out), str(out / src)], capture_output=True, check=True, timeout=90)
img.unlink()
print("\n".join(sorted(p.name for p in out.iterdir())))
